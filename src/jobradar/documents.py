"""Turn the small Markdown subset used for CVs and letters into DOCX and PDF. No photo, ever.

Supported: `# title`, `## section`, `- bullet`, `**bold**`, blank line = paragraph break.
DOCX uses python-docx; PDF uses reportlab with the built-in Helvetica (Latin-1 range: other
characters are folded to their closest ASCII form).
"""

from __future__ import annotations

import io
import re
import unicodedata
from xml.sax.saxutils import escape

_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _runs(text: str) -> list[tuple[str, bool]]:
    out, pos = [], 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out or [("", False)]


def _blocks(markdown: str) -> list[tuple[str, str]]:
    """(kind, text) with kind in title | heading | bullet | line | gap."""
    blocks = []
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            blocks.append(("gap", ""))
        elif line.startswith("# "):
            blocks.append(("title", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("heading", line[3:].strip()))
        elif line.lstrip().startswith(("- ", "* ")):
            blocks.append(("bullet", line.lstrip()[2:].strip()))
        else:
            blocks.append(("line", line.strip()))
    return blocks


def to_docx(markdown: str) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    for section in doc.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(1.8)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10.5)
    base.paragraph_format.space_after = Pt(2)
    for kind, text in _blocks(markdown):
        if kind == "gap":
            continue
        if kind == "bullet":
            para = doc.add_paragraph(style="List Bullet")
        else:
            para = doc.add_paragraph()
        if kind == "title":
            run = para.add_run(text)
            run.bold, run.font.size = True, Pt(18)
            continue
        if kind == "heading":
            run = para.add_run(text.upper())
            run.bold, run.font.size = True, Pt(11)
            para.paragraph_format.space_before = Pt(8)
            continue
        for chunk, bold in _runs(text):
            para.add_run(chunk).bold = bold
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pdf_safe(text: str) -> str:
    """Helvetica covers Latin-1 / Windows-1252. Anything else is folded to ASCII or replaced."""
    out = []
    for ch in text:
        try:
            ch.encode("cp1252")
            out.append(ch)
        except UnicodeEncodeError:
            folded = unicodedata.normalize("NFKD", ch).encode("ascii", "ignore").decode()
            out.append(folded or "?")
    return "".join(out)


def to_pdf(markdown: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=2)
    title = ParagraphStyle("title", parent=body, fontName="Helvetica-Bold", fontSize=18, leading=22, spaceAfter=4)
    heading = ParagraphStyle("heading", parent=body, fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=9,
                             spaceAfter=3)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=14, bulletIndent=3)

    def markup(text: str) -> str:
        return "".join(f"<b>{escape(pdf_safe(c))}</b>" if b else escape(pdf_safe(c)) for c, b in _runs(text))

    story = []
    for kind, text in _blocks(markdown):
        if kind == "gap":
            story.append(Spacer(1, 5))
        elif kind == "title":
            story.append(Paragraph(markup(text), title))
        elif kind == "heading":
            story.append(Paragraph(escape(pdf_safe(text.upper())), heading))
        elif kind == "bullet":
            story.append(Paragraph(markup(text), bullet, bulletText="•"))
        else:
            story.append(Paragraph(markup(text), body))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                      bottomMargin=1.8 * cm, title="", author="").build(story)
    return buf.getvalue()
